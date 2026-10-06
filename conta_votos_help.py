#!/usr/bin/env python3
"""Conta prints de votos para HELP em uma exportação ZIP do WhatsApp."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path, PurePosixPath

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MESSAGE_RE = re.compile(
    r"^[\ufeff\u200e\u200f\u202a-\u202e\u2066-\u2069]*"
    r"\[(?P<timestamp>\d{2}/\d{2}/\d{4}, \d{2}:\d{2}:\d{2})\] "
    r"(?P<sender>.*?): (?P<body>.*)$"
)
ATTACHMENT_RE = re.compile(r"<anexado:\s*(?P<name>[^>]+)>")
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}
HELP_RE = re.compile(r"\bhelp\b", re.IGNORECASE)
BMG_RE = re.compile(r"(?:\bbmg\b|bancobmg)", re.IGNORECASE)
VOTE_CONFIRMED_RE = re.compile(
    r"(?:agradece.{0,35}seu voto|voce\s+ja\s+votou\s+em)", re.IGNORECASE
)
VOTE_INVITATION_RE = re.compile(r"vote\s+agora(?:\s+\w{1,5}){0,2}\s+mesmo", re.IGNORECASE)


@dataclass
class ImageMessage:
    timestamp: str
    sender: str
    phone: str
    filename: str
    zip_entry: str | None


def clean_name(value: str) -> str:
    value = re.sub(r"[\u200e\u200f\u202a-\u202e\u2066-\u2069\ufeff]", "", value)
    value = value.replace("\u202f", " ").replace("\u00a0", " ").strip()
    if value.startswith("~"):
        value = value[1:].strip()
    return value


def sender_phone(sender: str) -> str:
    value = sender.replace("\u2010", "-").replace("\u2011", "-").replace("\u2013", "-")
    digits = sum(character.isdigit() for character in value)
    phone_characters_only = all(
        character.isdigit() or character.isspace() or character in "+-()." for character in value
    )
    if phone_characters_only and 10 <= digits <= 15:
        return value.strip()
    return "Não consta no ZIP"


def normalize_person_key(value: str) -> str:
    value = clean_name(value)
    value = unicodedata.normalize("NFKD", value).casefold()
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", value).split())


def read_people_dimension(
    path: Path,
) -> tuple[dict[str, dict[str, str]], dict[str, dict[str, str]]]:
    """Return the member dimension and a normalized WhatsApp-name lookup."""
    people: dict[str, dict[str, str]] = {}
    by_alias: dict[str, dict[str, str]] = {}
    if not path.is_file():
        raise ValueError(f"Dimensão de pessoas não encontrada: {path}")

    with path.open("r", encoding="utf-8-sig", newline="") as source:
        for row in csv.DictReader(source):
            name = (row.get("Nome no ranking") or "").strip()
            last4 = (row.get("Celular final 4") or "").strip()
            aliases = (row.get("Aliases WhatsApp") or "").split("|")
            if not name:
                continue
            if last4 and not re.fullmatch(r"\d{4}", last4):
                raise ValueError(f"O final do celular de {name} deve conter exatamente quatro dígitos.")
            person = {"Nome": name, "Celular final 4": last4}
            people[name] = person
            for alias in [name, *aliases]:
                key = normalize_person_key(alias)
                if not key:
                    continue
                existing = by_alias.get(key)
                if existing and existing["Nome"] != name:
                    raise ValueError(f"O alias '{alias}' aparece para mais de uma pessoa em {path.name}.")
                by_alias[key] = person
    return people, by_alias


def archive_index(archive: zipfile.ZipFile) -> tuple[dict[str, str], dict[str, list[str]]]:
    exact: dict[str, str] = {}
    by_basename: dict[str, list[str]] = defaultdict(list)
    for entry in archive.namelist():
        normalized = entry.replace("\\", "/").lstrip("./")
        exact[normalized.casefold()] = entry
        by_basename[PurePosixPath(normalized).name.casefold()].append(entry)
    return exact, by_basename


def find_entry(filename: str, exact: dict[str, str], by_basename: dict[str, list[str]]) -> str | None:
    normalized = filename.strip().replace("\\", "/").lstrip("./")
    match = exact.get(normalized.casefold())
    if match:
        return match
    candidates = by_basename.get(PurePosixPath(normalized).name.casefold(), [])
    return candidates[0] if len(candidates) == 1 else None


def read_messages(archive: zipfile.ZipFile) -> list[ImageMessage]:
    chat_entries = [n for n in archive.namelist() if PurePosixPath(n).name.casefold() == "_chat.txt"]
    if not chat_entries:
        raise ValueError("Não encontrei _chat.txt dentro do ZIP. Exporte a conversa com as mídias incluídas.")

    chat_entry = chat_entries[0]
    with archive.open(chat_entry) as stream:
        text = stream.read().decode("utf-8-sig", errors="replace")

    exact, by_basename = archive_index(archive)
    records: list[ImageMessage] = []
    for line in text.splitlines():
        message = MESSAGE_RE.match(line)
        if not message:
            continue
        timestamp = message.group("timestamp")
        sender = clean_name(message.group("sender"))
        body = message.group("body")
        for attachment in ATTACHMENT_RE.finditer(body):
            filename = attachment.group("name").strip()
            path = PurePosixPath(filename.replace("\\", "/"))
            if path.suffix.casefold() not in IMAGE_EXTENSIONS:
                continue
            # Figurinhas WEBP não são comprovantes de voto.
            if "-sticker-" in path.name.casefold():
                continue
            records.append(ImageMessage(
                timestamp=timestamp,
                sender=sender,
                phone=sender_phone(sender),
                filename=path.name,
                zip_entry=find_entry(filename, exact, by_basename),
            ))

    if not records:
        raise ValueError("Não encontrei anexos de imagem na conversa. Verifique se o ZIP inclui as mídias.")
    return records


def resolve_tesseract(configured: str) -> str:
    candidate = Path(configured)
    executable = str(candidate) if candidate.is_file() else shutil.which(configured)
    if not executable:
        raise RuntimeError(
            "Não encontrei o Tesseract. Instale o Tesseract OCR, adicione-o ao PATH "
            "ou rode novamente com --tesseract apontando para tesseract.exe."
        )

    result = subprocess.run(
        [executable, "--list-langs"], capture_output=True, text=True,
        encoding="utf-8", errors="replace", check=False,
    )
    languages = {line.strip() for line in result.stdout.splitlines() if line.strip()}
    missing = {"por", "eng"} - languages
    if result.returncode != 0 or missing:
        missing_text = ", ".join(sorted(missing)) if missing else "dados de idioma"
        raise RuntimeError(
            f"O Tesseract foi encontrado, mas faltam dados de idioma: {missing_text}. "
            "Instale os pacotes de idioma Portuguese (por) e English (eng)."
        )
    return executable


def classify(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text).casefold()
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    has_help = bool(HELP_RE.search(normalized))
    has_bmg = bool(BMG_RE.search(normalized))
    confirmed = bool(VOTE_CONFIRMED_RE.search(normalized))
    # O convite explícito sem frase de confirmação não prova que o voto foi enviado.
    if not confirmed and VOTE_INVITATION_RE.search(normalized):
        return "SEM COMPROVANTE"
    if has_help and has_bmg:
        return "REVISAR"
    if has_help:
        return "HELP"
    if has_bmg:
        return "BMG"
    return "REVISAR"


def review_reason(text: str, error: str, zip_entry: str | None) -> str:
    if zip_entry is None:
        return "Anexo de imagem não encontrado no ZIP."
    if error:
        return "Erro ao executar o OCR."
    normalized = unicodedata.normalize("NFKD", text).casefold()
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    if VOTE_INVITATION_RE.search(normalized) and not VOTE_CONFIRMED_RE.search(normalized):
        return "Convite para votar; a imagem não confirma que o voto foi enviado."
    has_help = bool(HELP_RE.search(normalized))
    has_bmg = bool(BMG_RE.search(normalized))
    if has_help and has_bmg:
        return "HELP e BMG apareceram juntos; confirmar qual empresa recebeu o voto."
    return "OCR não reconheceu HELP nem BMG; conferir o print."


def safe_console_print(message: str, stream) -> None:
    encoding = getattr(stream, "encoding", None) or "utf-8"
    printable = message.encode(encoding, errors="backslashreplace").decode(encoding, errors="replace")
    print(printable, file=stream)


def ocr_image(archive: zipfile.ZipFile, record: ImageMessage, executable: str) -> tuple[str, str]:
    if record.zip_entry is None:
        return "", "Arquivo de imagem não encontrado dentro do ZIP."

    suffix = PurePosixPath(record.filename).suffix or ".jpg"
    temp_path: str | None = None
    try:
        with archive.open(record.zip_entry) as source:
            with tempfile.NamedTemporaryFile(prefix="voto_", suffix=suffix, delete=False) as temp:
                temp_path = temp.name
                shutil.copyfileobj(source, temp)

        result = subprocess.run(
            [executable, temp_path, "stdout", "-l", "por+eng", "--psm", "6"],
            capture_output=True, text=True, encoding="utf-8", errors="replace", check=False,
        )
        if result.returncode != 0:
            error = result.stderr.strip() or f"Tesseract retornou código {result.returncode}."
            return result.stdout.strip(), error
        recognized = result.stdout.strip()
        if classify(recognized) != "REVISAR":
            return recognized, ""

        # Segunda tentativa apenas nos prints ambíguos: escala e aumenta contraste.
        try:
            from PIL import Image, ImageOps

            with Image.open(temp_path) as source_image:
                enhanced = ImageOps.autocontrast(ImageOps.grayscale(ImageOps.exif_transpose(source_image)))
                enhanced = enhanced.resize((enhanced.width * 2, enhanced.height * 2))
                with tempfile.NamedTemporaryFile(prefix="voto_ocr_", suffix=".png", delete=False) as temp:
                    enhanced_path = temp.name
                try:
                    enhanced.save(enhanced_path)
                    retry = subprocess.run(
                        [executable, enhanced_path, "stdout", "-l", "por+eng", "--psm", "6"],
                        capture_output=True, text=True, encoding="utf-8", errors="replace", check=False,
                    )
                    if retry.returncode == 0 and retry.stdout.strip():
                        recognized = f"{recognized}\n[OCR ampliado]\n{retry.stdout.strip()}"
                finally:
                    os.remove(enhanced_path)
        except (ImportError, OSError, ValueError):
            # A primeira leitura continua válida se o pré-processamento não estiver disponível.
            pass
        return recognized, ""
    except (OSError, zipfile.BadZipFile) as exc:
        return "", str(exc)
    finally:
        if temp_path:
            try:
                os.remove(temp_path)
            except OSError:
                pass


def set_sheet_style(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="183B56")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def copy_review_images(
    archive: zipfile.ZipFile,
    messages: list[ImageMessage],
    records: list[dict[str, str]],
    output_path: Path,
) -> Path | None:
    pending = [
        (message, record)
        for message, record in zip(messages, records)
        if record["Classificação"] == "REVISAR" and message.zip_entry is not None
    ]
    if not pending:
        return None

    output_path.parent.mkdir(parents=True, exist_ok=True)
    base = output_path.parent / f"{output_path.stem}_imagens_revisar"
    image_dir = base
    suffix = 2
    while image_dir.exists():
        image_dir = base.with_name(f"{base.name}_{suffix}")
        suffix += 1
    image_dir.mkdir()

    for index, (message, record) in enumerate(pending, start=1):
        image_path = image_dir / f"{index:03d}_{Path(message.filename).name}"
        with archive.open(message.zip_entry) as source, image_path.open("wb") as destination:
            shutil.copyfileobj(source, destination)
        record["Imagem local"] = str(image_path)
    return image_dir


def aggregate_people(
    records: list[dict[str, str]], people: dict[str, dict[str, str]]
) -> tuple[dict[str, Counter[str]], dict[str, str]]:
    totals: dict[str, Counter[str]] = defaultdict(Counter)
    last4_by_name = {name: person["Celular final 4"] for name, person in people.items()}
    for name in people:
        totals[name]
    for record in records:
        name = record["Pessoa"]
        totals[name][record["Classificação"]] += 1
        last4_by_name.setdefault(name, record["Celular final 4"])
    return totals, last4_by_name


def ordered_people(totals: dict[str, Counter[str]]) -> list[str]:
    return sorted(totals, key=lambda name: (-totals[name]["HELP"], name.casefold()))


def write_workbook(
    records: list[dict[str, str]],
    output_path: Path,
    people: dict[str, dict[str, str]],
) -> None:
    totals, last4_by_name = aggregate_people(records, people)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary.append([
        "Pessoa", "Celular (final 4)", "Votos HELP (OCR)", "BMG (não conta)",
        "Sem comprovante", "Revisar",
    ])
    for sender in ordered_people(totals):
        summary.append([
            sender,
            last4_by_name.get(sender, ""),
            totals[sender]["HELP"],
            totals[sender]["BMG"],
            totals[sender]["SEM COMPROVANTE"],
            totals[sender]["REVISAR"],
        ])
    summary.append([
        "TOTAL",
        "",
        sum(c["HELP"] for c in totals.values()),
        sum(c["BMG"] for c in totals.values()),
        sum(c["SEM COMPROVANTE"] for c in totals.values()),
        sum(c["REVISAR"] for c in totals.values()),
    ])
    set_sheet_style(summary, [30, 24, 20, 18, 20, 14])

    details = workbook.create_sheet("Detalhes")
    details.append([
        "Pessoa", "Remetente", "Celular (final 4)", "Data/hora", "Arquivo", "Classificação",
        "Motivo revisão", "Texto OCR", "Erro",
    ])
    for record in records:
        details.append([
            record["Pessoa"], record["Remetente"], record["Celular final 4"], record["Data/hora"],
            record["Arquivo"], record["Classificação"], record["Motivo revisão"],
            record["Texto OCR"], record["Erro"],
        ])
    set_sheet_style(details, [30, 30, 18, 22, 48, 18, 54, 72, 48])

    review = workbook.create_sheet("Revisar")
    review.append([
        "Pessoa", "Celular (final 4)", "Data/hora", "Arquivo", "Motivo da revisão",
        "Texto OCR", "Abrir print", "Erro",
    ])
    for record in records:
        if record["Classificação"] == "REVISAR" or record["Erro"]:
            review.append([
                record["Pessoa"], record["Celular final 4"], record["Data/hora"],
                record["Arquivo"], record["Motivo revisão"], record["Texto OCR"],
                "Abrir imagem" if record.get("Imagem local") else "Imagem indisponível",
                record["Erro"],
            ])
            link_cell = review.cell(review.max_row, 7)
            if record.get("Imagem local"):
                relative_path = os.path.relpath(record["Imagem local"], output_path.parent)
                link_cell.hyperlink = relative_path.replace(os.sep, "/")
                link_cell.style = "Hyperlink"
    set_sheet_style(review, [30, 24, 22, 48, 58, 72, 22, 48])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_ranking_csv(
    records: list[dict[str, str]],
    people: dict[str, dict[str, str]],
    output_path: Path,
    excluded_people: set[str] | None = None,
) -> None:
    totals, last4_by_name = aggregate_people(records, people)
    excluded_people = excluded_people or set()
    ranking_people = [
        name for name in ordered_people(totals)
        if normalize_person_key(name) not in excluded_people
    ]
    updated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as destination:
        writer = csv.writer(destination, delimiter=";")
        writer.writerow([
            "Posição", "Pessoa", "Celular final 4", "Votos HELP", "BMG (não conta)",
            "Sem comprovante", "Revisar", "Atualizado em",
        ])
        for position, name in enumerate(ranking_people, start=1):
            writer.writerow([
                position, name, last4_by_name.get(name, ""), totals[name]["HELP"],
                totals[name]["BMG"], totals[name]["SEM COMPROVANTE"],
                totals[name]["REVISAR"], updated_at,
            ])


def read_ocr_cache(path: Path | None) -> dict[tuple[str, str, str], str]:
    """Reuse prior OCR for matching sender/time/file entries in a previous report."""
    if path is None:
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Detalhes" not in workbook.sheetnames:
            raise ValueError("A planilha informada em --reuse-ocr-from não tem a aba Detalhes.")
        sheet = workbook["Detalhes"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1])}
        required = {"Remetente", "Data/hora", "Arquivo", "Texto OCR"}
        if not required.issubset(headers):
            raise ValueError("A aba Detalhes da planilha de OCR anterior não tem as colunas esperadas.")
        cache: dict[tuple[str, str, str], str] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            key = tuple(str(row[headers[name]] or "") for name in ("Remetente", "Data/hora", "Arquivo"))
            text = str(row[headers["Texto OCR"]] or "")
            if any(key) and text:
                cache[key] = text
        return cache
    finally:
        workbook.close()


def read_manual_corrections(path: Path) -> dict[str, str]:
    """Read visual-review overrides, keyed by the original WhatsApp filename."""
    if not path.is_file():
        return {}
    corrections: dict[str, str] = {}
    allowed = {"HELP", "BMG", "REVISAR", "SEM COMPROVANTE"}
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        for row in csv.DictReader(source):
            filename = Path((row.get("Arquivo") or "").strip()).name.casefold()
            label = (row.get("Classificação") or "").strip().upper()
            if not filename:
                continue
            if label not in allowed:
                raise ValueError(f"Classificação inválida para {filename} em {path.name}: {label}")
            corrections[filename] = label
    return corrections


def read_excluded_ranking_people(
    path: Path,
    people_by_alias: dict[str, dict[str, str]],
) -> set[str]:
    """Read local names/WhatsApp aliases that should not appear in the public ranking."""
    if not path.is_file():
        return set()

    excluded: set[str] = set()
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        if not reader.fieldnames or "Nome ou alias" not in reader.fieldnames:
            raise ValueError(f"{path.name} precisa ter a coluna 'Nome ou alias'.")
        for row in reader:
            value = (row.get("Nome ou alias") or "").strip()
            key = normalize_person_key(value)
            if not key:
                continue
            person = people_by_alias.get(key)
            excluded.add(normalize_person_key(person["Nome"]) if person else key)
    return excluded


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Conta prints de votos para HELP em uma exportação ZIP do WhatsApp."
    )
    parser.add_argument("--zip", required=True, type=Path, help="Arquivo ZIP exportado do grupo, com mídias.")
    parser.add_argument("--output", type=Path, help="Caminho opcional para o Excel de saída.")
    parser.add_argument(
        "--reuse-ocr-from", type=Path,
        help="Reaproveita o OCR de um Excel anterior para imagens com remetente, horário e arquivo iguais.",
    )
    parser.add_argument(
        "--tesseract", default="tesseract",
        help="Comando ou caminho completo do executável Tesseract (padrão: tesseract no PATH).",
    )
    args = parser.parse_args()

    if not args.zip.is_file():
        parser.error(f"Arquivo ZIP não encontrado: {args.zip}")
    if args.reuse_ocr_from and not args.reuse_ocr_from.is_file():
        parser.error(f"Planilha de OCR anterior não encontrada: {args.reuse_ocr_from}")
    output = args.output or args.zip.with_name(f"{args.zip.stem}_contagem_help.xlsx")

    try:
        executable = resolve_tesseract(args.tesseract)
        ocr_cache = read_ocr_cache(args.reuse_ocr_from)
        correction_path = Path(__file__).with_name("correcoes_votos.csv")
        manual_corrections = read_manual_corrections(correction_path)
        people_path = Path(__file__).with_name("dim_pessoas.csv")
        people, people_by_alias = read_people_dimension(people_path)
        exclusions_path = Path(__file__).with_name("pessoas_excluidas.csv")
        excluded_people = read_excluded_ranking_people(exclusions_path, people_by_alias)
        with zipfile.ZipFile(args.zip) as archive:
            messages = read_messages(archive)
            records: list[dict[str, str]] = []
            total = len(messages)
            for index, message in enumerate(messages, start=1):
                cache_key = (message.sender, message.timestamp, message.filename)
                cached_text = ocr_cache.get(cache_key)
                if cached_text and classify(cached_text) != "REVISAR":
                    text, error = cached_text, ""
                    safe_console_print(f"Reutilizando OCR {index}/{total}: {message.filename}", sys.stderr)
                else:
                    safe_console_print(f"Processando print {index}/{total}: {message.filename}", sys.stderr)
                    text, error = ocr_image(archive, message, executable)
                    if cached_text and cached_text not in text:
                        text = f"{cached_text}\n[OCR atualizado]\n{text}".strip()
                result = "REVISAR" if error else classify(text)
                correction = manual_corrections.get(message.filename.casefold())
                if correction:
                    result = correction
                person = people_by_alias.get(normalize_person_key(message.sender))
                display_name = person["Nome"] if person else clean_name(message.sender)
                last4 = person["Celular final 4"] if person else ""
                if not last4 and message.phone != "Não consta no ZIP":
                    last4 = "".join(character for character in message.phone if character.isdigit())[-4:]
                records.append({
                    "Pessoa": display_name,
                    "Remetente": clean_name(message.sender),
                    "Celular final 4": last4,
                    "Data/hora": message.timestamp,
                    "Arquivo": message.filename,
                    "Classificação": result,
                    "Motivo revisão": (
                        f"Correção manual registrada em {correction_path.name}." if correction
                        else review_reason(text, error, message.zip_entry)
                        if result in {"REVISAR", "SEM COMPROVANTE"} else ""
                    ),
                    "Texto OCR": text,
                    "Erro": error,
                })
            copy_review_images(archive, messages, records, output)
        write_workbook(records, output, people)
        ranking_output = output.with_name(f"{output.stem}_ranking.csv")
        write_ranking_csv(records, people, ranking_output, excluded_people)
    except (OSError, zipfile.BadZipFile, RuntimeError, ValueError) as exc:
        safe_console_print(f"Erro: {exc}", sys.stderr)
        return 1

    help_count = sum(record["Classificação"] == "HELP" for record in records)
    bmg_count = sum(record["Classificação"] == "BMG" for record in records)
    review_count = sum(record["Classificação"] == "REVISAR" for record in records)
    safe_console_print(f"Planilha criada: {output}", sys.stdout)
    safe_console_print(f"Ranking CSV criado: {ranking_output}", sys.stdout)
    no_proof_count = sum(record["Classificação"] == "SEM COMPROVANTE" for record in records)
    safe_console_print(
        f"HELP: {help_count} | BMG (não conta): {bmg_count} | "
        f"Sem comprovante: {no_proof_count} | Revisar: {review_count}",
        sys.stdout,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
